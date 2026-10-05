from datetime import date, datetime, timedelta, timezone
from io import BytesIO
import re
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import or_
from sqlalchemy.orm import Session

from .. import models, schemas, auth
from ..audit import record
from ..database import get_db

router = APIRouter(prefix="/time-entries", tags=["time-entries"])


@router.post("/start", response_model=schemas.TimeEntryOut)
def start_tracking(
    payload: schemas.TimeEntryStart,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_user),
):
    existing = (
        db.query(models.TimeEntry)
        .filter(
            models.TimeEntry.user_id == current_user.id,
            models.TimeEntry.status == models.TimeEntryStatus.active,
        )
        .first()
    )
    if existing:
        raise HTTPException(status_code=400, detail="A tracking session is already active")

    entry = models.TimeEntry(
        user_id=current_user.id,
        project_id=payload.project_id,
        start_ip_address=payload.ip_address,
        status=models.TimeEntryStatus.active,
        last_seen_at=datetime.now(timezone.utc),
        is_idle=False,
    )
    db.add(entry)
    db.commit()
    db.refresh(entry)
    return entry


@router.post("/{entry_id}/stop", response_model=schemas.TimeEntryOut)
def stop_tracking(
    entry_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_user),
):
    entry = (
        db.query(models.TimeEntry)
        .filter(models.TimeEntry.id == entry_id, models.TimeEntry.user_id == current_user.id)
        .first()
    )
    if not entry:
        raise HTTPException(status_code=404, detail="Time entry not found")
    if entry.status == models.TimeEntryStatus.stopped:
        return entry

    now = datetime.now(timezone.utc)
    entry.end_time = now
    entry.status = models.TimeEntryStatus.stopped
    start_time = entry.start_time
    if start_time.tzinfo is None:
        start_time = start_time.replace(tzinfo=timezone.utc)
    entry.duration_seconds = int((now - start_time).total_seconds())

    db.commit()
    db.refresh(entry)
    return entry


@router.get("/active", response_model=Optional[schemas.TimeEntryOut])
def get_active_entry(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_user),
):
    return (
        db.query(models.TimeEntry)
        .filter(
            models.TimeEntry.user_id == current_user.id,
            models.TimeEntry.status == models.TimeEntryStatus.active,
        )
        .first()
    )


@router.get("", response_model=List[schemas.TimeEntryOut])
def list_time_entries(
    user_id: Optional[int] = None,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_user),
):
    query = db.query(models.TimeEntry)
    if current_user.role in (models.UserRole.superadmin, models.UserRole.admin, models.UserRole.manager):
        if user_id is not None:
            query = query.filter(models.TimeEntry.user_id == user_id)
        visible_ids = [user.id for user in auth.visible_user_filter(db.query(models.User), current_user).all()]
        query = query.filter(models.TimeEntry.user_id.in_(visible_ids))
    else:
        visible_ids = auth.get_visible_member_ids(db, current_user)

    if user_id is not None and current_user.role != models.UserRole.superadmin:
        if not auth.can_manage(db, current_user, user_id):
            raise HTTPException(status_code=404, detail="User not found")

    if visible_ids is not None:
        query = query.filter(models.TimeEntry.user_id.in_(visible_ids))
    if user_id is not None:
        query = query.filter(models.TimeEntry.user_id == user_id)

    if start_date is not None:
        range_start = datetime.combine(start_date, datetime.min.time(), tzinfo=timezone.utc)
        query = query.filter(
            or_(models.TimeEntry.end_time.is_(None), models.TimeEntry.end_time >= range_start)
        )

    if end_date is not None:
        range_end = datetime.combine(end_date, datetime.max.time(), tzinfo=timezone.utc)
        query = query.filter(models.TimeEntry.start_time <= range_end)

    return query.order_by(models.TimeEntry.start_time.desc()).all()


def _format_duration(seconds):
    return f"{seconds // 3600}h {(seconds % 3600) // 60}m"


def _summarize_daily_entries(entries, screenshots_by_entry, now):
    entries_by_date = {}
    for entry in entries:
        entry_date = entry.start_time.date()
        entries_by_date.setdefault(entry_date, []).append(entry)

    daily_summaries = []
    for entry_date, day_entries in entries_by_date.items():
        first_login = None
        last_logout = None
        active_seconds = 0
        screenshot_count = 0
        activity_levels = []
        has_active_idle_session = False

        for entry in day_entries:
            start = entry.start_time
            if start.tzinfo is None:
                start = start.replace(tzinfo=timezone.utc)
            end = entry.end_time or (
                now if entry.status == models.TimeEntryStatus.active else start
            )
            if end.tzinfo is None:
                end = end.replace(tzinfo=timezone.utc)
            first_login = start if first_login is None else min(first_login, start)
            last_logout = end if last_logout is None else max(last_logout, end)
            active_seconds += (
                max(0, int((end - start).total_seconds()))
                if entry.status == models.TimeEntryStatus.active
                else max(0, entry.duration_seconds or 0)
            )

            activity = screenshots_by_entry.get(entry.id, [])
            screenshot_count += len(activity)
            activity_levels.extend(
                screenshot.activity_level
                for screenshot in activity
                if screenshot.activity_level is not None
            )
            has_active_idle_session = has_active_idle_session or (
                entry.status == models.TimeEntryStatus.active and entry.is_idle
            )

        span_seconds = max(0, int((last_logout - first_login).total_seconds()))
        idle_seconds = max(0, span_seconds - active_seconds)
        activity_details = f"{screenshot_count} screenshot(s)"
        if activity_levels:
            activity_details += f"; average activity {round(sum(activity_levels) / len(activity_levels))}%"
        if has_active_idle_session:
            activity_details += "; currently idle"

        daily_summaries.append({
            "date": entry_date,
            "first_login": first_login,
            "last_logout": last_logout,
            "session_duration_seconds": span_seconds,
            "active_seconds": active_seconds,
            "idle_seconds": idle_seconds,
            "session_count": len(day_entries),
            "activity_details": activity_details,
        })
    return daily_summaries


def _style_timesheet_worksheet(worksheet):
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1C2333")
    for column_cells in worksheet.columns:
        column_letter = column_cells[0].column_letter
        max_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 40)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions


@router.get("/report/download")
def download_company_daily_report(
    report_date: date,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_user),
):
    if current_user.role != models.UserRole.admin:
        raise HTTPException(status_code=403, detail="Only Admins can download company reports")

    company_member_ids = auth.get_descendant_ids(db, current_user.id)
    owned_department_ids = [
        department_id
        for (department_id,) in db.query(models.Department.id)
        .filter(models.Department.created_by_id == current_user.id)
        .all()
    ]
    if owned_department_ids:
        department_member_ids = {
            user_id
            for (user_id,) in db.query(models.User.id)
            .filter(models.User.department_id.in_(owned_department_ids))
            .all()
        }
        company_member_ids.update(department_member_ids)

    report_roles = (models.UserRole.manager, models.UserRole.tl, models.UserRole.user)
    report_users = (
        db.query(models.User)
        .filter(
            models.User.id.in_(company_member_ids),
            models.User.role.in_(report_roles),
        )
        .order_by(models.User.name.asc(), models.User.id.asc())
        .all()
    )
    if not report_users:
        raise HTTPException(status_code=404, detail="No report data found for the selected date")

    user_ids = [user.id for user in report_users]
    range_start = datetime.combine(report_date, datetime.min.time(), tzinfo=timezone.utc)
    range_end = datetime.combine(report_date + timedelta(days=1), datetime.min.time(), tzinfo=timezone.utc)
    entries = (
        db.query(models.TimeEntry)
        .filter(
            models.TimeEntry.user_id.in_(user_ids),
            models.TimeEntry.start_time >= range_start,
            models.TimeEntry.start_time < range_end,
        )
        .order_by(models.TimeEntry.start_time.asc())
        .all()
    )
    entry_ids = [entry.id for entry in entries]
    screenshots_by_entry = {}
    if entry_ids:
        for screenshot in db.query(models.Screenshot).filter(
            models.Screenshot.time_entry_id.in_(entry_ids)
        ).all():
            screenshots_by_entry.setdefault(screenshot.time_entry_id, []).append(screenshot)

    entries_by_user = {}
    for entry in entries:
        entries_by_user.setdefault(entry.user_id, []).append(entry)
    now = datetime.now(timezone.utc)

    def safe_cell(value):
        text_value = "" if value is None else str(value)
        if text_value.startswith(("=", "+", "-", "@", "\t", "\r")):
            return f"'{text_value}"
        return text_value

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Daily Report"
    worksheet.append([
        "User Name", "Email", "Department", "Role", "Date", "Login Time",
        "Logout Time", "Active Work Time", "Idle Time", "Number of Sessions",
        "Activity Details",
    ])
    for user in report_users:
        summaries = _summarize_daily_entries(
            entries_by_user.get(user.id, []),
            screenshots_by_entry,
            now,
        )
        summary = summaries[0] if summaries else None
        row = [
            safe_cell(user.name),
            safe_cell(user.email),
            safe_cell(user.department.name if user.department else ""),
            user.role.value,
            report_date,
        ]
        if summary:
            row.extend([
                summary["first_login"].time().replace(tzinfo=None),
                summary["last_logout"].time().replace(tzinfo=None),
                _format_duration(summary["active_seconds"]),
                _format_duration(summary["idle_seconds"]),
                summary["session_count"],
                safe_cell(summary["activity_details"]),
            ])
        else:
            row.extend([None] * 6)
        worksheet.append(row)
        row_number = worksheet.max_row
        worksheet.cell(row_number, 5).number_format = "dd-mmm-yyyy"
        worksheet.cell(row_number, 6).number_format = "hh:mm:ss"
        worksheet.cell(row_number, 7).number_format = "hh:mm:ss"

    _style_timesheet_worksheet(worksheet)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    record(
        db,
        current_user,
        "timesheet.company_report_downloaded",
        "company",
        details={"report_date": report_date.isoformat(), "member_count": len(report_users)},
    )
    db.commit()
    filename = f"company_timesheet_{report_date.isoformat()}.xlsx"
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/{user_id}/download")
def download_timesheet(
    user_id: int,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_user),
):
    if current_user.role not in (models.UserRole.superadmin, models.UserRole.admin):
        raise HTTPException(
            status_code=403,
            detail="Only Admins and Super Admins can download timesheets",
        )

    if start_date is not None and end_date is not None and start_date > end_date:
        raise HTTPException(status_code=400, detail="Start date must be on or before end date")

    target_user = db.query(models.User).filter(models.User.id == user_id).first()
    if not target_user:
        raise HTTPException(status_code=404, detail="User not found")

    if current_user.role == models.UserRole.admin and current_user.id != target_user.id:
        allowed_roles = {models.UserRole.manager, models.UserRole.tl, models.UserRole.user}
        if (
            target_user.role not in allowed_roles
            or not auth.can_manage(db, current_user, target_user.id)
        ):
            raise HTTPException(status_code=403, detail="You cannot download this user's timesheet")

    query = db.query(models.TimeEntry).filter(models.TimeEntry.user_id == target_user.id)
    if start_date is not None:
        range_start = datetime.combine(start_date, datetime.min.time(), tzinfo=timezone.utc)
        query = query.filter(
            or_(models.TimeEntry.end_time.is_(None), models.TimeEntry.end_time >= range_start)
        )
    if end_date is not None:
        range_end = datetime.combine(end_date, datetime.max.time(), tzinfo=timezone.utc)
        query = query.filter(models.TimeEntry.start_time <= range_end)

    entries = query.order_by(models.TimeEntry.start_time.asc()).all()
    if not entries:
        raise HTTPException(status_code=404, detail="No timesheet data found for the selected date range")

    entry_ids = [entry.id for entry in entries]
    screenshots = (
        db.query(models.Screenshot)
        .filter(
            models.Screenshot.user_id == target_user.id,
            models.Screenshot.time_entry_id.in_(entry_ids),
        )
        .all()
    )
    screenshots_by_entry = {}
    for screenshot in screenshots:
        screenshots_by_entry.setdefault(screenshot.time_entry_id, []).append(screenshot)

    daily_summaries = _summarize_daily_entries(
        entries,
        screenshots_by_entry,
        datetime.now(timezone.utc),
    )

    def csv_value(value):
        text_value = "" if value is None else str(value)
        if text_value.startswith(("=", "+", "-", "@", "\t", "\r")):
            return f"'{text_value}"
        return text_value

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Timesheet"
    worksheet.append([
        "User Name", "Email", "Department", "Role", "Date",
        "Login Time (UTC)", "Logout Time (UTC)",
        "Active Work Time", "Idle Time", "Number of Sessions", "Activity Details",
    ])
    for summary in daily_summaries:
        row = [
            csv_value(target_user.name),
            csv_value(target_user.email),
            csv_value(target_user.department.name if target_user.department else ""),
            target_user.role.value,
            summary["date"],
            summary["first_login"].time().replace(tzinfo=None),
            summary["last_logout"].time().replace(tzinfo=None),
            _format_duration(summary["active_seconds"]),
            _format_duration(summary["idle_seconds"]),
            summary["session_count"],
            csv_value(summary["activity_details"]),
        ]
        worksheet.append(row)
        row_number = worksheet.max_row
        worksheet.cell(row_number, 5).number_format = "dd-mmm-yyyy"
        worksheet.cell(row_number, 6).number_format = "hh:mm:ss"
        worksheet.cell(row_number, 7).number_format = "hh:mm:ss"

    safe_name = re.sub(r"[^A-Za-z0-9_-]+", "_", target_user.name).strip("_") or f"user_{user_id}"
    _style_timesheet_worksheet(worksheet)

    output = BytesIO()
    workbook.save(output)
    filename = f"timesheet_{safe_name}_{start_date or 'all'}_to_{end_date or 'all'}.xlsx"
    record(
        db,
        current_user,
        "timesheet.downloaded",
        "user",
        target_user.id,
        {"start_date": str(start_date) if start_date else None, "end_date": str(end_date) if end_date else None},
    )
    db.commit()
    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/{entry_id}/heartbeat", response_model=schemas.TimeEntryOut)
def heartbeat(
    entry_id: int,
    payload: schemas.TimeEntryHeartbeat,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_user),
):
    entry = (
        db.query(models.TimeEntry)
        .filter(
            models.TimeEntry.id == entry_id,
            models.TimeEntry.user_id == current_user.id,
            models.TimeEntry.status == models.TimeEntryStatus.active,
        )
        .first()
    )
    if not entry:
        raise HTTPException(status_code=404, detail="Active time entry not found")

    entry.last_seen_at = datetime.now(timezone.utc)
    entry.is_idle = payload.is_idle
    db.commit()
    db.refresh(entry)
    return entry