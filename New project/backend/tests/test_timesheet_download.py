import asyncio
from io import BytesIO
import unittest
from datetime import datetime, timedelta, timezone

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app import models
from app.routers.time_entries import download_company_daily_report, download_timesheet


class TimesheetDownloadTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite:///:memory:")
        models.Base.metadata.create_all(self.engine)
        self.db = sessionmaker(bind=self.engine)()

        self.department = models.Department(name=f"Department {id(self)}")
        self.db.add(self.department)
        self.db.flush()

        def user(name, role, parent=None, department=True):
            member = models.User(
                name=name,
                email=f"{name.lower().replace(' ', '-')}-{id(self)}@test",
                hashed_password="x",
                role=role,
                parent=parent,
                department_id=self.department.id if department else None,
            )
            self.db.add(member)
            self.db.flush()
            self.db.add(models.TimeEntry(
                user_id=member.id,
                start_time=datetime.now(timezone.utc) - timedelta(hours=2),
                end_time=datetime.now(timezone.utc) - timedelta(hours=1),
                duration_seconds=3600,
                status=models.TimeEntryStatus.stopped,
            ))
            return member

        self.superadmin = user("Super Admin", models.UserRole.superadmin, department=False)
        self.admin = user("Admin", models.UserRole.admin, self.superadmin)
        self.admin_peer = user("Admin Peer", models.UserRole.admin, self.superadmin)
        self.manager = user("Manager", models.UserRole.manager, self.admin)
        self.team_lead = user("Team Lead", models.UserRole.tl, self.manager)
        self.member = user("Member", models.UserRole.user, self.team_lead)
        self.other_manager = user("Other Manager", models.UserRole.manager, self.admin_peer)
        self.other_lead = user("Other Lead", models.UserRole.tl, self.other_manager)
        self.other_member = user("Other Member", models.UserRole.user, self.other_lead)
        self.db.commit()

    def tearDown(self):
        self.db.close()
        self.engine.dispose()

    def download(self, actor, target):
        return download_timesheet(
            user_id=target.id,
            start_date=None,
            end_date=None,
            db=self.db,
            current_user=actor,
        )

    def company_report(self, actor, report_date):
        return download_company_daily_report(
            report_date=report_date,
            db=self.db,
            current_user=actor,
        )

    @staticmethod
    def is_xlsx(response):
        return response.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def test_superadmin_can_download_every_role(self):
        for target in (
            self.superadmin, self.admin, self.manager, self.team_lead, self.member,
        ):
            with self.subTest(target=target.role):
                self.assertTrue(self.is_xlsx(self.download(self.superadmin, target)))

    def test_admin_can_download_lower_roles_but_not_admin_peers_or_superadmin(self):
        for target in (self.admin, self.manager, self.team_lead, self.member):
            with self.subTest(target=target.role):
                self.assertTrue(self.is_xlsx(self.download(self.admin, target)))
        for target in (self.admin_peer, self.superadmin):
            with self.subTest(target=target.role):
                with self.assertRaises(HTTPException) as raised:
                    self.download(self.admin, target)
                self.assertEqual(raised.exception.status_code, 403)

    def test_manager_cannot_download_any_timesheets(self):
        for target in (self.manager, self.team_lead, self.member):
            with self.subTest(target=target.role, target_id=target.id):
                with self.assertRaises(HTTPException) as raised:
                    self.download(self.manager, target)
                self.assertEqual(raised.exception.status_code, 403)

    def test_team_lead_cannot_download_any_timesheets(self):
        for target in (self.team_lead, self.member, self.manager):
            with self.subTest(target=target.role, target_id=target.id):
                with self.assertRaises(HTTPException) as raised:
                    self.download(self.team_lead, target)
                self.assertEqual(raised.exception.status_code, 403)

    def test_user_cannot_download_any_timesheets(self):
        for target in (self.member, self.other_member):
            with self.subTest(target_id=target.id):
                with self.assertRaises(HTTPException) as raised:
                    self.download(self.member, target)
                self.assertEqual(raised.exception.status_code, 403)

    def test_admin_can_download_own_but_not_other_admin_timesheets(self):
        self.assertTrue(self.is_xlsx(self.download(self.admin, self.admin)))
        for target in (self.admin_peer, self.superadmin):
            with self.assertRaises(HTTPException) as raised:
                self.download(self.admin, target)
            self.assertEqual(raised.exception.status_code, 403)

    def test_report_includes_session_activity_headers_and_audit_event(self):
        entry = self.db.query(models.TimeEntry).filter_by(user_id=self.member.id).first()
        second_start = entry.end_time + timedelta(minutes=20)
        self.db.add(models.TimeEntry(
            user_id=self.member.id,
            start_time=second_start,
            end_time=second_start + timedelta(minutes=10),
            duration_seconds=600,
            status=models.TimeEntryStatus.stopped,
        ))
        third_start = entry.start_time + timedelta(days=1)
        self.db.add(models.TimeEntry(
            user_id=self.member.id,
            start_time=third_start,
            end_time=third_start + timedelta(minutes=5),
            duration_seconds=300,
            status=models.TimeEntryStatus.stopped,
        ))
        self.db.add(models.Screenshot(
            time_entry_id=entry.id,
            user_id=self.member.id,
            file_path="unused",
            ip_address="127.0.0.1",
            activity_level=80,
        ))
        self.db.commit()

        response = self.download(self.admin, self.member)
        content = asyncio.run(self._read_body(response))
        workbook = load_workbook(BytesIO(content), data_only=True)
        worksheet = workbook.active
        headers = [cell.value for cell in worksheet[1]]
        row = dict(zip(headers, [cell.value for cell in worksheet[2]]))

        self.assertIn("attachment; filename=\"timesheet_Member_all_to_all.xlsx\"", response.headers["content-disposition"])
        self.assertTrue(self.is_xlsx(response))
        self.assertEqual(worksheet.max_row, 3)
        self.assertEqual(worksheet["E2"].value.strftime("%d-%b-%Y"), entry.start_time.strftime("%d-%b-%Y"))
        self.assertEqual(worksheet["E2"].number_format, "dd-mmm-yyyy")
        self.assertEqual(worksheet["F2"].value.strftime("%H:%M:%S"), entry.start_time.strftime("%H:%M:%S"))
        self.assertEqual(
            worksheet["G2"].value.strftime("%H:%M:%S"),
            (second_start + timedelta(minutes=10)).strftime("%H:%M:%S"),
        )
        self.assertNotIn("Total Tracked Hours", row)
        self.assertNotIn("Status", row)
        self.assertNotIn("Session Duration", row)
        self.assertEqual(row["Active Work Time"], "1h 10m")
        self.assertEqual(row["Idle Time"], "0h 20m")
        self.assertEqual(row["Number of Sessions"], 2)
        self.assertIn("average activity 80%", row["Activity Details"])
        self.assertEqual(worksheet["E3"].value.date(), third_start.date())
        self.assertEqual(worksheet["F3"].value.strftime("%H:%M:%S"), third_start.strftime("%H:%M:%S"))
        self.assertEqual(worksheet["G3"].value.strftime("%H:%M:%S"), (third_start + timedelta(minutes=5)).strftime("%H:%M:%S"))
        self.assertEqual(worksheet["H3"].value, "0h 5m")
        self.assertEqual(worksheet["I3"].value, "0h 0m")
        self.assertEqual(worksheet["J3"].value, 1)
        event = self.db.query(models.AuditEvent).filter_by(action="timesheet.downloaded").one()
        self.assertEqual(event.actor_id, self.admin.id)
        self.assertEqual(event.target_id, str(self.member.id))

    def test_admin_company_report_aggregates_selected_date_across_departments(self):
        outside_manager = models.User(
            name="Outside Manager",
            email=f"outside-manager-{id(self)}@test",
            hashed_password="x",
            role=models.UserRole.manager,
            parent=self.admin,
            department_id=None,
        )
        self.db.add(outside_manager)
        self.db.flush()
        self.db.add(models.TimeEntry(
            user_id=outside_manager.id,
            start_time=datetime.now(timezone.utc) - timedelta(hours=2),
            end_time=datetime.now(timezone.utc) - timedelta(hours=1),
            duration_seconds=3600,
            status=models.TimeEntryStatus.stopped,
        ))
        manager_entry = self.db.query(models.TimeEntry).filter_by(user_id=self.manager.id).first()
        second_start = manager_entry.end_time + timedelta(minutes=20)
        self.db.add(models.TimeEntry(
            user_id=self.manager.id,
            start_time=second_start,
            end_time=second_start + timedelta(minutes=10),
            duration_seconds=600,
            status=models.TimeEntryStatus.stopped,
        ))
        next_day_start = manager_entry.start_time + timedelta(days=1)
        self.db.add(models.TimeEntry(
            user_id=self.member.id,
            start_time=next_day_start,
            end_time=next_day_start + timedelta(minutes=5),
            duration_seconds=300,
            status=models.TimeEntryStatus.stopped,
        ))
        empty_member = models.User(
            name="Empty Member",
            email=f"report-empty-{id(self)}@test",
            hashed_password="x",
            role=models.UserRole.user,
            parent=self.team_lead,
            department_id=self.department.id,
        )
        self.db.add(empty_member)
        self.db.commit()

        response = self.company_report(self.admin, manager_entry.start_time.date())
        content = asyncio.run(self._read_body(response))
        worksheet = load_workbook(BytesIO(content), data_only=True).active
        headers = [cell.value for cell in worksheet[1]]
        rows = [dict(zip(headers, [cell.value for cell in row])) for row in worksheet.iter_rows(min_row=2)]
        rows_by_name = {row["User Name"]: row for row in rows}

        self.assertTrue(self.is_xlsx(response))
        self.assertEqual(
            headers,
            [
                "User Name", "Email", "Department", "Role", "Date", "Login Time",
                "Logout Time", "Active Work Time", "Idle Time", "Number of Sessions",
                "Activity Details",
            ],
        )
        self.assertEqual(set(rows_by_name), {
            "Manager", "Team Lead", "Member", "Empty Member", "Outside Manager",
        })
        self.assertNotIn("Admin", rows_by_name)
        self.assertNotIn("Admin Peer", rows_by_name)
        self.assertNotIn("Other Manager", rows_by_name)
        self.assertNotIn("Other Lead", rows_by_name)
        self.assertNotIn("Other Member", rows_by_name)
        self.assertEqual(rows_by_name["Outside Manager"]["Number of Sessions"], 1)
        manager_row = rows_by_name["Manager"]
        self.assertEqual(manager_row["Role"], "manager")
        self.assertEqual(manager_row["Department"], self.department.name)
        self.assertEqual(manager_row["Date"].date(), manager_entry.start_time.date())
        self.assertEqual(manager_row["Login Time"].strftime("%H:%M:%S"), manager_entry.start_time.strftime("%H:%M:%S"))
        self.assertEqual(manager_row["Logout Time"].strftime("%H:%M:%S"), (second_start + timedelta(minutes=10)).strftime("%H:%M:%S"))
        self.assertEqual(manager_row["Active Work Time"], "1h 10m")
        self.assertEqual(manager_row["Idle Time"], "0h 20m")
        self.assertEqual(manager_row["Number of Sessions"], 2)
        self.assertTrue(all(
            rows_by_name["Empty Member"][header] is None
            for header in headers[5:]
        ))
        self.assertEqual(rows_by_name["Empty Member"]["Department"], self.department.name)
        self.assertEqual(rows_by_name["Empty Member"]["Role"], "user")
        self.assertEqual(rows_by_name["Empty Member"]["Date"].date(), manager_entry.start_time.date())
        event = self.db.query(models.AuditEvent).filter_by(
            action="timesheet.company_report_downloaded"
        ).one()
        self.assertEqual(event.actor_id, self.admin.id)
        self.assertEqual(event.details["report_date"], manager_entry.start_time.date().isoformat())

    def test_company_report_is_admin_only_and_includes_company_roster_without_entries(self):
        report_date = (datetime.now(timezone.utc) + timedelta(days=10)).date()
        for actor in (self.superadmin, self.manager, self.team_lead, self.member):
            with self.subTest(actor=actor.role):
                with self.assertRaises(HTTPException) as denied:
                    self.company_report(actor, report_date)
                self.assertEqual(denied.exception.status_code, 403)

        response = self.company_report(self.admin, report_date)
        content = asyncio.run(self._read_body(response))
        worksheet = load_workbook(BytesIO(content), data_only=True).active
        headers = [cell.value for cell in worksheet[1]]
        rows = [dict(zip(headers, [cell.value for cell in row])) for row in worksheet.iter_rows(min_row=2)]

        self.assertTrue(self.is_xlsx(response))
        self.assertEqual(len(rows), 3)
        self.assertEqual(
            {row["User Name"] for row in rows},
            {"Manager", "Team Lead", "Member"},
        )
        for row in rows:
            with self.subTest(user=row["User Name"]):
                self.assertTrue(row["User Name"])
                self.assertTrue(row["Email"])
                self.assertEqual(row["Department"], self.department.name)
                self.assertIn(row["Role"], {"manager", "tl", "user"})
                self.assertEqual(row["Date"].date(), report_date)
                self.assertTrue(all(row[header] is None for header in headers[5:]))

    def test_invalid_user_date_range_and_empty_report_are_rejected(self):
        with self.assertRaises(HTTPException) as missing:
            download_timesheet(99999, None, None, self.db, self.superadmin)
        self.assertEqual(missing.exception.status_code, 404)

        with self.assertRaises(HTTPException) as invalid_range:
            download_timesheet(
                self.member.id,
                datetime(2026, 10, 5).date(),
                datetime(2026, 10, 4).date(),
                self.db,
                self.admin,
            )
        self.assertEqual(invalid_range.exception.status_code, 400)

        no_entries = models.User(
            name="Empty User",
            email=f"empty-{id(self)}@test",
            hashed_password="x",
            role=models.UserRole.user,
            parent=self.team_lead,
            department_id=self.department.id,
        )
        self.db.add(no_entries)
        self.db.commit()
        with self.assertRaises(HTTPException) as empty:
            self.download(self.admin, no_entries)
        self.assertEqual(empty.exception.status_code, 404)
        self.assertIn("No timesheet data", empty.exception.detail)

    @staticmethod
    async def _read_body(response):
        chunks = [chunk async for chunk in response.body_iterator]
        return b"".join(chunk.encode() if isinstance(chunk, str) else chunk for chunk in chunks)


if __name__ == "__main__":
    unittest.main()
